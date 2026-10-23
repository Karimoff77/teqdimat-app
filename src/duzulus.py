"""
DUZULUS.PY — slayd tipləri üçün alternativ düzülüşlər (eyni məzmun, fərqli struktur).
qur.py hər təqdimatda bunları avtomatik qarışdırır; "duzulus" sahəsi ilə əl ilə də seçilə bilər.
"""
import math
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.dml.color import RGBColor

C, M = PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE

# tip -> mümkün düzülüşlər (ilki mühərrikdəki ilkin düzülüşdür)
VARIANTS = {
    "maddeler":   ["list", "numbered", "split", "cards2"],
    "kartlar":    ["grid", "rows", "icons"],
    "muqayise":   ["cards", "versus"],
    "proses":     ["horizontal", "vertical", "chevron"],
    "statistika": ["plain", "cards"],
    "sekilli":    ["right", "left", "top"],
    "plan":       ["list", "cards", "line"],
    "tesekkur":   ["center", "split", "question"],
    "cedvel":     ["table", "cards"],
}


def on(c):
    """Rəngin üstündə oxunaqlı mətn rəngi."""
    return RGBColor(0x14, 0x14, 0x14) if (0.299*c[0] + 0.587*c[1] + 0.114*c[2]) > 150 else RGBColor(0xFF, 0xFF, 0xFF)


def box(sl, kind, x, y, w, h, fill=None, line=None, lw=1.0):
    sp = sl.shapes.add_shape(kind, int(x), int(y), int(w), int(h))
    if fill is None: sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def circ(d, sl, cx, cy, diam, text, size=15, fill=None):
    fill = fill or d.t["accent"]
    box(sl, MSO_SHAPE.OVAL, cx - diam // 2, cy - diam // 2, diam, diam, fill=fill)
    d.textbox(sl, text, cx - diam // 2, cy - diam // 2, diam, diam, size=size, bold=True,
              color=on(fill), align=C, anchor=M)


def head(d, sl, s, chk, kicker_key="kicker"):
    d.label(sl, chk("v.kicker", s[kicker_key], 30))
    d.title(sl, chk("v.baslik", s["baslik"], 48))


# ============================================================ MADDELER
def maddeler(d, s, v, chk):
    items = [chk("v.meqam", m, 80) for m in s["meqamlar"][:6]]
    t = d.t
    if v == "numbered":
        sl = d.new_slide(); head(d, sl, s, chk)
        y, h, gap = Inches(1.95), Inches(0.68), Inches(0.1)
        for i, m in enumerate(items):
            d.card(sl, Inches(0.6), y, Inches(12.1), h)
            circ(d, sl, Inches(1.2), y + h // 2, Inches(0.46), str(i + 1))
            d.textbox(sl, m, Inches(1.75), y, Inches(10.8), h, size=17, color=t["text_main"], anchor=M)
            y += h + gap
    elif v == "split":
        sl = d.new_slide(footer=False)
        box(sl, MSO_SHAPE.RECTANGLE, 0, 0, Inches(4.6), d.prs.slide_height, fill=t["accent"])
        d.textbox(sl, s["kicker"].upper(), Inches(0.6), Inches(0.9), Inches(3.6), Inches(0.4),
                  size=13, bold=True, color=on(t["accent"]))
        d.textbox(sl, chk("v.baslik", s["baslik"], 48), Inches(0.6), Inches(1.5), Inches(3.6), Inches(4.5),
                  size=32, bold=True, color=on(t["accent"]), font=t["font_head"])
        n = len(items); step = min(0.95, 5.6 / n); y = Inches(1.05)
        for i, m in enumerate(items):
            d.circle_dot(sl, Inches(5.45), y + Inches(0.4), Inches(0.16))
            d.textbox(sl, m, Inches(5.8), y, Inches(6.9), Inches(0.85), size=19, color=t["text_main"], anchor=M)
            if i < n - 1:
                d.divider(sl, y + Inches(step - 0.06), x=Inches(5.3), w=Inches(7.4))
            y += Inches(step)
    else:  # cards2
        sl = d.new_slide(); head(d, sl, s, chk)
        n = len(items); rows = math.ceil(n / 2)
        cw, ch, gap = Inches(5.95), Inches(1.35), Inches(0.2)
        for i, m in enumerate(items):
            x = Inches(0.6) + (i % 2) * (cw + gap); y = Inches(2.0) + (i // 2) * (ch + gap)
            d.card(sl, x, y, cw, ch)
            box(sl, MSO_SHAPE.RECTANGLE, x + Inches(0.3), y + Inches(0.3), Inches(0.5), Inches(0.06), fill=t["accent"])
            d.textbox(sl, f"{i+1:02d}", x + Inches(0.3), y + Inches(0.4), Inches(0.9), Inches(0.5), size=22, bold=True,
                      color=t["accent"], font=t["font_head"])
            d.textbox(sl, m, x + Inches(1.3), y + Inches(0.1), cw - Inches(1.55), ch - Inches(0.2), size=17,
                      color=t["text_main"], anchor=M)
    return sl


# ============================================================ KARTLAR
def kartlar(d, s, v, chk):
    items = [(chk("v.kb", a, 30), chk("v.kz", b, 80)) for a, b in s["kartlar"][:6]]
    t = d.t
    sl = d.new_slide(); head(d, sl, s, chk)
    if v == "rows":
        y, h, gap = Inches(1.95), Inches(0.66), Inches(0.1)
        for i, (a, b) in enumerate(items):
            d.card(sl, Inches(0.6), y, Inches(12.1), h)
            d.textbox(sl, f"{i+1:02d}", Inches(0.85), y, Inches(0.9), h, size=22, bold=True, color=t["accent"],
                      font=t["font_head"], anchor=M)
            d.textbox(sl, a, Inches(1.8), y, Inches(3.7), h, size=16, bold=True, color=t["text_main"], anchor=M)
            d.textbox(sl, b, Inches(5.7), y, Inches(6.8), h, size=14, color=t["text_body"], anchor=M)
            y += h + gap
    else:  # icons
        n = len(items); cols = 2 if n == 4 else 3; gap = Inches(0.3)
        cw = int((Inches(12.1) - gap * (cols - 1)) / cols); ch = Inches(2.15)
        for i, (a, b) in enumerate(items):
            x = Inches(0.6) + (i % cols) * (cw + gap); y = Inches(1.95) + (i // cols) * (ch + gap)
            d.card(sl, x, y, cw, ch)
            circ(d, sl, x + Inches(0.65), y + Inches(0.65), Inches(0.62), str(i + 1), size=17)
            d.textbox(sl, a, x + Inches(1.15), y + Inches(0.3), cw - Inches(1.35), Inches(0.7), size=15, bold=True,
                      color=t["text_main"], anchor=M)
            d.textbox(sl, b, x + Inches(0.3), y + Inches(1.2), cw - Inches(0.6), Inches(0.85), size=13,
                      color=t["text_body"])
    return sl


# ============================================================ MÜQAYİSƏ
def muqayise(d, s, v, chk):
    t = d.t
    sl = d.new_slide(); head(d, sl, s, chk)
    cw, y, h = Inches(5.5), Inches(1.95), Inches(4.5)
    sides = [(Inches(0.6), chk("v.sb", s["sol_baslik"], 34), s["sol"][:4], t["accent"]),
             (Inches(7.23), chk("v.gb", s["sag_baslik"], 34), s["sag"][:4], t["accent2"])]
    for x, hd, bl, col in sides:
        d.card(sl, x, y, cw, h)
        box(sl, MSO_SHAPE.RECTANGLE, x + Inches(0.3), y + Inches(0.3), cw - Inches(0.6), Inches(0.65), fill=col)
        d.textbox(sl, hd, x + Inches(0.3), y + Inches(0.3), cw - Inches(0.6), Inches(0.65), size=18, bold=True,
                  color=on(col), align=C, anchor=M, font=t["font_head"])
        yy = y + Inches(1.3)
        for m in bl:
            chk("v.m", m, 80)
            d.circle_dot(sl, x + Inches(0.5), yy + Inches(0.3), Inches(0.14), color=t["accent"])
            d.textbox(sl, m, x + Inches(0.8), yy, cw - Inches(1.1), Inches(0.65), size=16, color=t["text_body"], anchor=M)
            yy += Inches(0.78)
    circ(d, sl, Inches(6.665), y + h // 2, Inches(0.85), "VS", size=15, fill=t["text_main"])
    if s.get("izah"):
        d.textbox(sl, s["izah"], Inches(0.6), Inches(6.55), Inches(12.1), Inches(0.45), size=15, bold=True,
                  color=t["accent"], align=C)
    return sl


# ============================================================ PROSES
def proses(d, s, v, chk):
    t = d.t
    steps = [(chk("v.ab", a, 36), chk("v.az", b, 90)) for a, b in s["addimlar"][:4]]
    n = len(steps)
    sl = d.new_slide(); head(d, sl, s, chk)
    if v == "vertical":
        sp = 1.12 if n == 4 else 1.35
        y0 = Inches(2.05)
        box(sl, MSO_SHAPE.RECTANGLE, Inches(1.33), y0 + Inches(0.25), Inches(0.05), Inches(sp * (n - 1)), fill=t["card_line"])
        for i, (a, b) in enumerate(steps):
            y = y0 + Inches(sp * i)
            circ(d, sl, Inches(1.355), y + Inches(0.27), Inches(0.55), str(i + 1), size=16)
            d.textbox(sl, a, Inches(2.1), y - Inches(0.03), Inches(10.5), Inches(0.4), size=18, bold=True, color=t["text_main"])
            d.textbox(sl, b, Inches(2.1), y + Inches(0.36), Inches(10.5), Inches(0.55), size=14, color=t["text_body"])
        if s.get("netice"):
            d.card(sl, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.58))
            d.textbox(sl, s["netice"], Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.58), size=16, bold=True,
                      color=t["accent"], align=C, anchor=M)
    else:  # chevron
        cw = int(Inches(12.1) / n)
        for i, (a, b) in enumerate(steps):
            x = Inches(0.6) + cw * i
            kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
            box(sl, kind, x, Inches(2.2), cw, Inches(1.05), fill=t["accent"], line=t["bg"], lw=3)
            d.textbox(sl, a, x + Inches(0.45), Inches(2.2), cw - Inches(0.85), Inches(1.05), size=15, bold=True,
                      color=on(t["accent"]), align=C, anchor=M)
            d.textbox(sl, b, x + Inches(0.15), Inches(3.55), cw - Inches(0.3), Inches(2.0), size=14,
                      color=t["text_body"], align=C)
        if s.get("netice"):
            d.card(sl, Inches(0.6), Inches(5.9), Inches(12.1), Inches(0.75))
            d.textbox(sl, s["netice"], Inches(0.6), Inches(5.9), Inches(12.1), Inches(0.75), size=18, bold=True,
                      color=t["accent"], align=C, anchor=M)
    return sl


# ============================================================ STATİSTİKA
def statistika(d, s, v, chk):
    t = d.t
    rows = s["rəqəmlər"][:4]; n = len(rows)
    mx = {2: 9, 3: 6, 4: 5}.get(n, 9)
    sl = d.new_slide(); head(d, sl, s, chk)
    gap = Inches(0.3); cw = int((Inches(12.1) - gap * (n - 1)) / n)
    for i, (num, lab) in enumerate(rows):
        x = Inches(0.6) + i * (cw + gap); y = Inches(2.3)
        d.card(sl, x, y, cw, Inches(3.1))
        box(sl, MSO_SHAPE.RECTANGLE, x + int(cw / 2) - Inches(0.4), y + Inches(0.3), Inches(0.8), Inches(0.07), fill=t["accent"])
        d.textbox(sl, chk("v.rq", num, mx), x, y + Inches(0.55), cw, Inches(1.2), size=54, bold=True, color=t["accent"],
                  font=t["font_head"], align=C, anchor=M)
        d.textbox(sl, chk("v.iz", lab, 60), x + Inches(0.2), y + Inches(1.9), cw - Inches(0.4), Inches(1.0), size=16,
                  color=t["text_body"], align=C)
    return sl


# ============================================================ ŞƏKİLLİ
def sekilli(d, s, v, chk):
    """Mətn + şəkil yeri (PowerPoint-in real şəkil placeholder-i). v: right | left | top"""
    t = d.t
    d.blank = d.prs.slide_layouts[8]
    sl = d.new_slide()
    d.blank = d.prs.slide_layouts[6]
    pic = None
    for ph in list(sl.placeholders):
        if "PICTURE" in str(ph.placeholder_format.type):
            pic = ph
        else:
            ph._element.getparent().remove(ph._element)
    head(d, sl, s, chk)
    items = [chk("v.meqam", m, 75) for m in s["meqamlar"][:4]]
    if v == "top":
        X, Y, W, H = Inches(0.6), Inches(1.75), Inches(12.1), Inches(2.75)
        n = len(items); gap = Inches(0.25); cw = int((Inches(12.1) - gap * (n - 1)) / n)
        for i, m in enumerate(items):
            x = Inches(0.6) + i * (cw + gap)
            d.card(sl, x, Inches(4.75), cw, Inches(1.95))
            box(sl, MSO_SHAPE.RECTANGLE, x + Inches(0.25), Inches(4.98), Inches(0.5), Inches(0.06), fill=t["accent"])
            d.textbox(sl, m, x + Inches(0.25), Inches(5.15), cw - Inches(0.5), Inches(1.45), size=15, color=t["text_main"])
        lab_y, hint_y = Y + Inches(0.8), Y + Inches(1.3)
    else:
        bx = Inches(0.85) if v == "right" else Inches(6.75)
        X = Inches(7.2) if v == "right" else Inches(0.6)
        Y, W, H = Inches(1.9), Inches(5.5), Inches(4.4)
        y = Inches(2.1)
        for m in items:
            d.circle_dot(sl, bx, y + Inches(0.17), Inches(0.12))
            d.textbox(sl, m, bx + Inches(0.25), y, Inches(5.5), Inches(0.85), size=17, color=t["text_body"])
            y += Inches(0.95)
        lab_y, hint_y = Y + Inches(1.55), Y + Inches(2.1)
    card = d.card(sl, X, Y, W, H)
    card.line.color.rgb = t["accent"]; card.line.width = Pt(1.5); card.line.dash_style = MSO_LINE.DASH
    ip = s.get("sekil_axtaris", "")
    d.textbox(sl, "ŞƏKİL YERİ", X, lab_y, W, Inches(0.5), size=16, bold=True, color=t["accent"], align=C)
    d.textbox(sl, "Şəkil əlavə etmək üçün kliklə" + (f"\nAxtarış: {ip}" if ip else ""), X + Inches(0.3), hint_y,
              W - Inches(0.6), Inches(1.0), size=13, color=t["text_body"], align=C)
    if pic is not None:
        pic.left, pic.top, pic.width, pic.height = X + Inches(0.1), Y + Inches(0.1), W - Inches(0.2), H - Inches(0.2)
        sl.shapes._spTree.append(pic._element)
    return sl


# ============================================================ PLAN
def plan(d, s, v, chk):
    t = d.t
    items = [(chk("v.q", a, 24), chk("v.i", b, 70)) for a, b in s["maddeler"][:7]]
    n = len(items)
    sl = d.new_slide(); d.label(sl, "PLAN"); d.title(sl, "Təqdimatın Planı")
    if v == "line" and n <= 5:
        cw = int(Inches(12.1) / n); y = Inches(2.6)
        box(sl, MSO_SHAPE.RECTANGLE, Inches(0.6) + cw // 2, y + Inches(0.25), cw * (n - 1), Inches(0.05), fill=t["card_line"])
        for i, (a, b) in enumerate(items):
            x = Inches(0.6) + cw * i
            circ(d, sl, x + cw // 2, y + Inches(0.275), Inches(0.6), str(i + 1), size=17)
            d.textbox(sl, a, x + Inches(0.1), y + Inches(0.95), cw - Inches(0.2), Inches(0.8), size=16, bold=True,
                      color=t["text_main"], align=C)
            d.textbox(sl, b, x + Inches(0.1), y + Inches(1.75), cw - Inches(0.2), Inches(1.7), size=13,
                      color=t["text_body"], align=C)
    else:
        cols = 2 if n <= 4 else (3 if n <= 6 else 4)
        rows = math.ceil(n / cols); gap = Inches(0.25)
        cw = int((Inches(12.1) - gap * (cols - 1)) / cols); ch = Inches(2.15) if cols > 2 else Inches(2.0)
        for i, (a, b) in enumerate(items):
            x = Inches(0.6) + (i % cols) * (cw + gap); y = Inches(2.0) + (i // cols) * (ch + gap)
            d.card(sl, x, y, cw, ch)
            d.textbox(sl, f"{i+1:02d}", x + Inches(0.3), y + Inches(0.2), Inches(1.2), Inches(0.5), size=22, bold=True,
                      color=t["accent"], font=t["font_head"])
            d.textbox(sl, a, x + Inches(0.3), y + Inches(0.75), cw - Inches(0.6), Inches(0.6), size=16, bold=True,
                      color=t["text_main"])
            d.textbox(sl, b, x + Inches(0.3), y + Inches(1.35), cw - Inches(0.6), Inches(0.75), size=13, color=t["text_body"])
    return sl


# ============================================================ TƏŞƏKKÜR
def tesekkur(d, s, v, chk):
    t = d.t; f = s.get("alt", ""); W, H = d.prs.slide_width, d.prs.slide_height
    sl = d.new_slide(footer=False)
    if v == "split":
        box(sl, MSO_SHAPE.RECTANGLE, 0, 0, Inches(5.4), H, fill=t["accent"])
        d.textbox(sl, "TƏŞƏKKÜR\nEDİRİK", Inches(0.6), Inches(2.5), Inches(4.4), Inches(2.2), size=40, bold=True,
                  color=on(t["accent"]), font=t["font_head"])
        d.textbox(sl, "Diqqətinizə görə təşəkkür edirik.", Inches(6.2), Inches(2.4), Inches(6.5), Inches(0.9), size=24,
                  color=t["text_main"])
        if f:
            d.card(sl, Inches(6.2), Inches(3.7), Inches(6.4), Inches(1.0))
            d.textbox(sl, f, Inches(6.2), Inches(3.7), Inches(6.4), Inches(1.0), size=18, bold=True, color=t["accent"],
                      align=C, anchor=M, font=t["font_head"])
        d.textbox(sl, "Suallarınızı gözləyirəm.", Inches(6.2), Inches(5.1), Inches(6.5), Inches(0.5), size=15,
                  color=t["text_body"])
    else:  # question
        d.textbox(sl, "Suallarınız?", Inches(0.8), Inches(2.2), W - Inches(1.6), Inches(1.6), size=56, bold=True,
                  color=t["text_main"], align=C, font=t["font_head"])
        box(sl, MSO_SHAPE.RECTANGLE, int((W - Inches(1.4)) / 2), Inches(4.0), Inches(1.4), Inches(0.07), fill=t["accent"])
        d.textbox(sl, "Diqqətinizə görə təşəkkür edirik", Inches(0.8), Inches(4.3), W - Inches(1.6), Inches(0.6), size=20,
                  color=t["text_body"], align=C)
        if f:
            d.textbox(sl, f, Inches(0.8), Inches(5.2), W - Inches(1.6), Inches(0.6), size=18, bold=True,
                      color=t["accent"], align=C, font=t["font_head"])
    return sl


# ============================================================ CƏDVƏL
def cedvel(d, s, v, chk):
    rows = [[chk("v.cell", c, 48) for c in r] for r in s["setirler"][:6]]
    heads = s["basliqlar"]; t = d.t
    if len(rows) > 5 or len(heads) != 3:
        return d.cedvel_slaydi(chk("v.kicker", s["kicker"], 30), chk("v.baslik", s["baslik"], 48), heads, rows)
    sl = d.new_slide(); head(d, sl, s, chk)
    xs = [Inches(0.85), Inches(4.5), Inches(8.4)]; ws = [Inches(3.4), Inches(3.7), Inches(4.1)]
    for j, hd in enumerate(heads):
        d.textbox(sl, hd.upper(), xs[j], Inches(1.85), ws[j], Inches(0.35), size=12, bold=True, color=t["accent"])
    y, h, gap = Inches(2.25), Inches(0.74), Inches(0.1)
    for r in rows:
        d.card(sl, Inches(0.6), y, Inches(12.1), h)
        d.textbox(sl, r[0], xs[0], y, ws[0], h, size=16, bold=True, color=t["text_main"], anchor=M)
        d.textbox(sl, r[1], xs[1], y, ws[1], h, size=14, color=t["text_body"], anchor=M)
        d.textbox(sl, r[2], xs[2], y, ws[2], h, size=14, color=t["text_body"], anchor=M)
        y += h + gap
    return sl
