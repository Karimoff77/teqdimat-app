"""
QUR.PY — JSON -> .pptx (fənndən asılı olmayan renderer)
İstifadə:  python qur.py decks/iqtisadiyyat.json
JSON-u gələcəkdə LLM yaradacaq; burada hər sahənin uzunluq limiti yoxlanılır.
"""
import json, sys
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
import random
from engine import TeqdimatDeck
import duzulus
from dogrula import dogrula

WARN = []

def chk(yer, metn, mx, setir_ucun=False):
    """Uzunluq limiti. setir_ucun=True olduqda hər sətir ayrıca yoxlanılır."""
    parts = metn.split("\n") if setir_ucun else [metn]
    for p in parts:
        if len(p) > mx:
            WARN.append(f"{yer}: {len(p)} simvol (limit {mx}) -> '{p[:35]}…'")
    return metn

def lum(c):
    return (0.299*c[0] + 0.587*c[1] + 0.114*c[2])

def seriya_rengi(d, i):
    bg = d.t["bg"]
    for cand in (d.t["accent"], d.t["accent2"], d.t["text_main"]):
        if abs(lum(cand) - lum(bg)) > 90 and (i == 0 or cand != d.t["accent"]):
            pass
    pal = [d.t["accent"], d.t["accent2"], d.t["text_body"]]
    c = pal[i % len(pal)]
    if abs(lum(c) - lum(bg)) < 70:      # arxa fonla zəif kontrast
        c = d.t["text_main"]
    return c

def qrafik(d, s):
    sl = d.new_slide()
    d.label(sl, chk("qrafik.kicker", s["kicker"], 30))
    d.title(sl, chk("qrafik.baslik", s["baslik"], 48))
    cd = CategoryChartData()
    cd.categories = s["kateqoriyalar"]
    for ser in s["setirler"]:
        cd.add_series(ser["ad"], ser["qiymetler"])
    tip = XL_CHART_TYPE.LINE_MARKERS if s.get("novu", "xett") == "xett" else XL_CHART_TYPE.COLUMN_CLUSTERED
    neticeler = s.get("neticeler", [])
    cw = Inches(7.6) if neticeler else Inches(12.1)
    d.card(sl, Inches(0.6), Inches(1.9), cw, Inches(4.9))
    gf = sl.shapes.add_chart(tip, Inches(0.8), Inches(2.0), cw - Inches(0.4), Inches(4.7), cd)
    ch = gf.chart
    ch.has_title = False
    ch.font.size = Pt(13); ch.font.color.rgb = d.t["text_body"]
    ch.has_legend = len(s["setirler"]) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(13); ch.legend.font.color.rgb = d.t["text_main"]
    va, ca = ch.value_axis, ch.category_axis
    va.major_gridlines.format.line.color.rgb = d.t["card_line"]
    va.format.line.fill.background()
    ca.format.line.color.rgb = d.t["card_line"]
    for ax, ad in ((ca, s.get("x_oxu")), (va, s.get("y_oxu"))):
        if ad:
            ax.has_title = True
            ax.axis_title.text_frame.text = ad
            r = ax.axis_title.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(12); r.font.bold = False; r.font.color.rgb = d.t["text_body"]
    for i, ser in enumerate(ch.plots[0].series):
        c = seriya_rengi(d, i)
        if tip == XL_CHART_TYPE.COLUMN_CLUSTERED:
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = c
        else:
            ser.format.line.color.rgb = c; ser.format.line.width = Pt(3); ser.smooth = False
            ser.marker.size = 9
            ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = c
            ser.marker.format.line.color.rgb = c
    if neticeler:
        x = Inches(8.45); y = Inches(2.1)
        d.textbox(sl, "ƏSAS FİKİR", x, y, Inches(4.2), Inches(0.4), size=13, bold=True, color=d.t["accent"])
        y += Inches(0.6)
        for n in neticeler[:3]:
            chk("qrafik.netice", n, 80)
            d.circle_dot(sl, x + Inches(0.1), y + Inches(0.17), Inches(0.12))
            d.textbox(sl, n, x + Inches(0.3), y, Inches(3.9), Inches(1.2), size=15, color=d.t["text_main"])
            y += Inches(1.3)

def build(spec):
    import engine as _eng
    _seed = str(spec.get("seed", spec["movzu"]))
    _r = random.Random("stil|" + _seed)
    tema = spec.get("tema", "corporate")
    if tema == "auto":
        tema = _r.choice(sorted(_eng.THEMES))
    d = TeqdimatDeck(theme_name=tema, movzu_adi=spec["movzu"], muellif=spec.get("muellif", ""))
    if spec.get("stil", "auto") == "auto":      # dekor, künc, başlıq və titul müstəqil qarışdırılır
        d.st.update(decor=_r.choice(["none", "left_bar", "top_bar", "corner_circle", "blocks", "frame"]),
                    radius=_r.choice([0.0, 0.04, 0.1, 0.18, 0.25]),
                    header=_r.choice(["plain", "tick", "band", "pill"]),
                    titul=_r.choice(["left", "center", "band"]))
    rng = random.Random(str(spec.get("seed", spec["movzu"])))
    last = {}
    def pick(tip, s):
        opts = duzulus.VARIANTS.get(tip)
        if not opts: return None
        v = s.get("duzulus")
        if v not in opts:
            cand = [o for o in opts if o != last.get(tip)] or opts
            v = rng.choices(cand, weights=[1 if o == opts[0] else 2 for o in cand])[0]   # yeni düzülüşlərə üstünlük
        last[tip] = v
        return v
    for i, s in enumerate(spec["slaydlar"], 1):
        t, ctx = s["tip"], f"s{i}.{s['tip']}"
        V = pick(t, s)
        if V is not None and V != duzulus.VARIANTS[t][0]:
            getattr(duzulus, t)(d, s, V, lambda yer, m, mx, st=False, _c=ctx: chk(_c + "." + yer, m, mx, st))
            continue
        if t == "titul":
            if s.get("stil_titul") in ("left", "center", "band"):   # Manual rejim: titul düzülüşünü slayd özü seçir
                d.st["titul"] = s["stil_titul"]
            d.titul_slaydi(chk(ctx+".kicker", s["kicker"], 40),
                           chk(ctx+".baslik", s["baslik"], 33, True),
                           chk(ctx+".alt", s["alt_baslik"], 110),
                           [chk(ctx+".info", x, 60) for x in s.get("info", [])],
                           s.get("tarix_yer", ""))
        elif t == "plan":
            d.plan_slaydi("Təqdimatın Planı",
                          [(f"{k:02d}", chk(ctx+".qisa", a, 24), chk(ctx+".izah", b, 70))
                           for k, (a, b) in enumerate(s["maddeler"][:7], 1)])
        elif t == "maddeler":
            chk(ctx+".sayi", "x"*len(s["meqamlar"]), 6)
            d.movzu_slaydi(chk(ctx+".kicker", s["kicker"], 30), chk(ctx+".baslik", s["baslik"], 48),
                           [chk(ctx+".meqam", m, 80) for m in s["meqamlar"]])
        elif t == "kartlar":
            d.grid_slaydi(chk(ctx+".kicker", s["kicker"], 30), chk(ctx+".baslik", s["baslik"], 48),
                          [(f"{k:02d}", chk(ctx+".kb", a, 30), chk(ctx+".kz", b, 80))
                           for k, (a, b) in enumerate(s["kartlar"][:6], 1)], cols=3)
        elif t == "muqayise":
            d.iki_kart_slaydi(chk(ctx+".kicker", s["kicker"], 30), chk(ctx+".baslik", s["baslik"], 48),
                              chk(ctx+".sb", s["sol_baslik"], 34), [chk(ctx+".sm", m, 80) for m in s["sol"][:4]],
                              chk(ctx+".gb", s["sag_baslik"], 34), [chk(ctx+".gm", m, 80) for m in s["sag"][:4]],
                              s.get("izah", ""))
        elif t == "proses":
            d.proses_axini_slaydi(chk(ctx+".kicker", s["kicker"], 30), chk(ctx+".baslik", s["baslik"], 48),
                                  [(f"{k:02d}", chk(ctx+".ab", a, 36), chk(ctx+".az", b, 90))
                                   for k, (a, b) in enumerate(s["addimlar"][:4], 1)],
                                  s.get("netice", ""))
        elif t == "xett":
            d.timeline_slaydi(chk(ctx+".kicker", s["kicker"], 30), chk(ctx+".baslik", s["baslik"], 48),
                              [tuple(x) for x in s["hadiseler"][:4]], s.get("esas_fikir", ""))
        elif t == "statistika":
            d.stat_slaydi(chk(ctx+".kicker", s["kicker"], 30), chk(ctx+".baslik", s["baslik"], 48),
                          [(chk(ctx+".rq", a, {1:12,2:9,3:6,4:5}[min(len(s["rəqəmlər"]),4)]), chk(ctx+".iz", b, 60)) for a, b in s["rəqəmlər"][:4]])
        elif t == "cedvel":
            d.cedvel_slaydi(chk(ctx+".kicker", s["kicker"], 30), chk(ctx+".baslik", s["baslik"], 48),
                            s["basliqlar"], [[chk(ctx+".cell", c, 48) for c in r] for r in s["setirler"][:6]])
        elif t == "iddia":
            d.iddia_slaydi(chk(ctx+".kicker", s["kicker"], 30), chk(ctx+".baslik", s["baslik"], 48),
                           chk(ctx+".iddia", s["iddia"], 80), chk(ctx+".izah", s.get("izah", ""), 150),
                           s.get("simvol", ""))
        elif t == "formula":
            d.formula_slaydi(chk(ctx+".kicker", s["kicker"], 30), chk(ctx+".baslik", s["baslik"], 48),
                             s["formula"], [tuple(x) for x in s["aciqlamalar"][:3]])
        elif t == "sitat":
            d.sitat_slaydi(chk(ctx+".sitat", s["sitat"], 140), s.get("muellif", ""))
        elif t == "qrafik":
            qrafik(d, s)
        elif t == "sekilli":
            duzulus.sekilli(d, s, "right", lambda yer, m, mx, st=False, _c=ctx: chk(_c + "." + yer, m, mx, st))
        elif t == "tesekkur":
            d.teşekkur_slaydi(s.get("alt", ""))
        else:
            WARN.append(f"{ctx}: naməlum slayd tipi")
    return d

if __name__ == "__main__":
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    spec, REPAIR = dogrula(spec)
    if REPAIR:
        print(f"  ~ {len(REPAIR)} avtomatik düzəliş:")
        for r in REPAIR: print("   ~", r)
    d = build(spec)
    out = sys.argv[2] if len(sys.argv) > 2 else spec["fayl"]
    d.saxla(out)
    print(out, "|", len(d.prs.slides), "slayd |", len(WARN), "xəbərdarlıq")
    for w in WARN: print("  !", w)
