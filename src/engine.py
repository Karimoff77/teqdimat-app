"""
ENGINE.PY — Təqdimat Mühərriki
Bu fayla normalda toxunmaq lazım deyil.
Yeni slayd tipi lazım olanda mən sizə yalnız 
əlavə ediləcək funksiyanı verəcəyəm.
"""

import math
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree


# ============================================================
# TEMALAR
# ============================================================
THEMES = {
 "emerald": {
        "bg":         RGBColor(0x0D, 0x1B, 0x16),
        "card":       RGBColor(0x15, 0x26, 0x1F),
        "card_line":  RGBColor(0x23, 0x40, 0x2F),
        "text_main":  RGBColor(0xF2, 0xF7, 0xF2),
        "text_body":  RGBColor(0xA9, 0xC2, 0xAE),
        "text_dim":   RGBColor(0x5C, 0x72, 0x67),
        "accent":     RGBColor(0xE8, 0xB8, 0x4B),
        "accent2":    RGBColor(0xFF, 0x7A, 0x59),
        "font_head":  "Georgia",
        "font_body":  "Calibri",
    },
    "indigo": {
        "bg":         RGBColor(0x14, 0x10, 0x24),
        "card":       RGBColor(0x1F, 0x1A, 0x38),
        "card_line":  RGBColor(0x35, 0x2C, 0x5C),
        "text_main":  RGBColor(0xF5, 0xF3, 0xFF),
        "text_body":  RGBColor(0xB8, 0xAE, 0xD9),
        "text_dim":   RGBColor(0x6C, 0x62, 0x94),
        "accent":     RGBColor(0xFF, 0xD1, 0x66),   # Qızılı
        "accent2":    RGBColor(0x06, 0xD6, 0xA0),   # Mint-yaşıl
        "font_head":  "Georgia",
        "font_body":  "Calibri",
    },
    "corporate": {
        "bg": RGBColor(0xF7, 0xF8, 0xFA), "card": RGBColor(0xFF, 0xFF, 0xFF),
        "card_line": RGBColor(0xE2, 0xE5, 0xEA), "text_main": RGBColor(0x1A, 0x1F, 0x2B),
        "text_body": RGBColor(0x5B, 0x63, 0x70), "text_dim": RGBColor(0xA0, 0xA6, 0xB0),
        "accent": RGBColor(0x25, 0x63, 0xEB), "accent2": RGBColor(0xDB, 0xEA, 0xFE),
        "font_head": "Georgia", "font_body": "Calibri",
    },
    "sunset": {
        "bg": RGBColor(0x2B, 0x17, 0x15), "card": RGBColor(0x3D, 0x22, 0x1E),
        "card_line": RGBColor(0x5A, 0x35, 0x2C), "text_main": RGBColor(0xFD, 0xF1, 0xE8),
        "text_body": RGBColor(0xCF, 0xA8, 0x94), "text_dim": RGBColor(0x8C, 0x6B, 0x5C),
        "accent": RGBColor(0xF2, 0xA0, 0x4D), "accent2": RGBColor(0xE8, 0x5D, 0x4A),
        "font_head": "Georgia", "font_body": "Calibri",
    },
    "ocean": {
        "bg": RGBColor(0x0A, 0x1A, 0x2E), "card": RGBColor(0x12, 0x28, 0x42),
        "card_line": RGBColor(0x1F, 0x3B, 0x5C), "text_main": RGBColor(0xE8, 0xF1, 0xFB),
        "text_body": RGBColor(0x8F, 0xAC, 0xC9), "text_dim": RGBColor(0x4F, 0x66, 0x80),
        "accent": RGBColor(0x4F, 0xD1, 0xE8), "accent2": RGBColor(0xA7, 0x8B, 0xFA),
        "font_head": "Trebuchet MS", "font_body": "Calibri",
    },
    "mono": {
        "bg": RGBColor(0x00, 0x00, 0x00), "card": RGBColor(0x18, 0x18, 0x18),
        "card_line": RGBColor(0x33, 0x33, 0x33), "text_main": RGBColor(0xFF, 0xFF, 0xFF),
        "text_body": RGBColor(0xAA, 0xAA, 0xAA), "text_dim": RGBColor(0x66, 0x66, 0x66),
        "accent": RGBColor(0xFF, 0xFF, 0xFF), "accent2": RGBColor(0xFF, 0xD6, 0x00),
        "font_head": "Arial Black", "font_body": "Arial",
    },
    "pastel": {
        "bg": RGBColor(0xFA, 0xF3, 0xE8), "card": RGBColor(0xFF, 0xFF, 0xFF),
        "card_line": RGBColor(0xEA, 0xDD, 0xC8), "text_main": RGBColor(0x3D, 0x32, 0x2A),
        "text_body": RGBColor(0x8A, 0x7A, 0x68), "text_dim": RGBColor(0xC2, 0xB3, 0x9E),
        "accent": RGBColor(0xE0, 0x7A, 0x5F), "accent2": RGBColor(0x7D, 0x9D, 0x8A),
        "font_head": "Georgia", "font_body": "Calibri",
    },
}


def _T(bg, card, line, main, body, dim, acc, acc2, fh, fb):
    h = lambda x: RGBColor.from_string(x)
    return {"bg": h(bg), "card": h(card), "card_line": h(line), "text_main": h(main),
            "text_body": h(body), "text_dim": h(dim), "accent": h(acc), "accent2": h(acc2),
            "font_head": fh, "font_body": fb}

THEMES.update({
    "lavender": _T("F5F0FF","FFFFFF","E4DAFB","2B1D4A","6B5B8E","B2A6CE","7C3AED","F472B6","Trebuchet MS","Calibri"),
    "forest":   _T("F1F5EE","FFFFFF","D3E0CB","1F3A2A","5C7566","A3B5A8","2F7D4F","F2B84B","Cambria","Calibri"),
    "rose":     _T("FFF5F5","FFFFFF","FBD5D5","4A1D24","8C5A62","D1A9AE","E11D48","FB923C","Georgia","Calibri"),
    "graphite": _T("1C1F26","262A33","3A3F4B","F3F4F6","AEB4C0","6B7280","38BDF8","F472B6","Arial","Arial"),
    "sand":     _T("F4EEE3","FBF8F1","DDD2BD","2F2A22","7A6E5C","BDB19C","8C5A2B","3F7F7A","Palatino Linotype","Calibri"),
    "neon":     _T("0B0B14","15152A","2A2A50","FFFFFF","B6B6E0","6A6A9C","FF3EA5","00E5FF","Century Gothic","Calibri"),
    "sky":      _T("EEF6FF","FFFFFF","CFE3FA","0F2A47","4F6B8A","A3B8CF","0EA5E9","F59E0B","Trebuchet MS","Calibri"),
    "academic": _T("FFFFFF","F5F5F4","D6D3D1","1C1917","57534E","A8A29E","7F1D1D","B45309","Cambria","Calibri"),
})

# Üslub: dekor növü və kartın künc yumşaqlığı (rəngdən asılı olmayan "şablon" fərqi)
STYLE_DEFAULT = {"decor": "none", "radius": 0.07, "titul": "left", "header": "plain"}
STYLES = {
    "emerald":  {"decor": "left_bar",      "radius": 0.07},
    "indigo":   {"decor": "corner_circle", "radius": 0.12, "titul": "center"},
    "corporate":{"decor": "top_bar",       "radius": 0.04, "titul": "band"},
    "sunset":   {"decor": "blocks",        "radius": 0.10, "titul": "band"},
    "ocean":    {"decor": "frame",         "radius": 0.06, "titul": "band"},
    "mono":     {"decor": "none",          "radius": 0.0},
    "pastel":   {"decor": "corner_circle", "radius": 0.18, "titul": "center"},
    "lavender": {"decor": "corner_circle", "radius": 0.22, "titul": "center"},
    "forest":   {"decor": "left_bar",      "radius": 0.03},
    "rose":     {"decor": "frame",         "radius": 0.14, "titul": "center"},
    "graphite": {"decor": "top_bar",       "radius": 0.0},
    "sand":     {"decor": "blocks",        "radius": 0.02},
    "neon":     {"decor": "corner_circle", "radius": 0.25, "titul": "center"},
    "sky":      {"decor": "top_bar",       "radius": 0.16, "titul": "band"},
    "academic": {"decor": "frame",         "radius": 0.0, "titul": "band"},
}

def _on(c):
    return RGBColor(0x14, 0x14, 0x14) if (0.299*c[0] + 0.587*c[1] + 0.114*c[2]) > 150 else RGBColor(0xFF, 0xFF, 0xFF)


def set_transition(slide, effect="fade", speed="med"):
    sld = slide._element
    existing = sld.find(qn('p:transition'))
    if existing is not None:
        sld.remove(existing)
    if effect == "none":
        return
    ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    xml_str = f'<p:transition xmlns:p="{ns}" spd="{speed}"><p:{effect}/></p:transition>'
    el = etree.fromstring(xml_str.encode('utf-8'))
    cSld = sld.find(qn('p:cSld'))
    cSld.addnext(el)


# ============================================================
# DECK SİNİFİ
# ============================================================
class TeqdimatDeck:

    def __init__(self, theme_name="ocean", movzu_adi="Təqdimat", muellif=""):
        if theme_name not in THEMES:
            raise ValueError(f"Naməlum tema: {theme_name}. Mövcud: {list(THEMES.keys())}")
        self.t = dict(THEMES[theme_name])
        self.st = {**STYLE_DEFAULT, **STYLES.get(theme_name, {})}
        self.theme_name = theme_name
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.33)
        self.prs.slide_height = Inches(7.5)
        self.blank = self.prs.slide_layouts[6]
        self.movzu_adi = movzu_adi
        self.muellif = muellif
        self.slide_no = 0

    # ---------- TƏMƏL ----------

    def new_slide(self, footer=True, transition="fade"):
        slide = self.prs.slides.add_slide(self.blank)
        self.slide_no += 1
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.t["bg"]
        self._decor(slide)
        if footer:
            self._header_deco(slide)
        if footer:
            self._footer(slide)
            self._page_number(slide)
        set_transition(slide, transition, "med")
        return slide

    def textbox(self, slide, text, x, y, w, h, size=18, bold=False,
                color=None, align=PP_ALIGN.LEFT, font=None, anchor=MSO_ANCHOR.TOP):
        color = color or self.t["text_main"]
        font = font or self.t["font_body"]
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = color
            p.font.name = font
            p.alignment = align
        return box

    def label(self, slide, text, x=Inches(0.6), y=Inches(0.5)):
        if self.st.get("header") == "pill":
            tt = text.upper(); w = Inches(0.5 + 0.105 * len(tt))
            pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.34))
            try: pill.adjustments[0] = 0.5
            except Exception: pass
            pill.fill.solid(); pill.fill.fore_color.rgb = self.t["accent"]
            pill.line.fill.background(); pill.shadow.inherit = False
            return self.textbox(slide, tt, x, y, w, Inches(0.34), size=12, bold=True, color=_on(self.t["accent"]),
                                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return self.textbox(slide, text.upper(), x, y, Inches(8), Inches(0.4),
                             size=13, bold=True, color=self.t["accent"])

    def title(self, slide, text, x=Inches(0.6), y=Inches(0.75), w=Inches(11.5), size=32):
        if self.st.get("header") == "pill" and y == Inches(0.75):
            y = Inches(0.88)
        return self.textbox(slide, text, x, y, w, Inches(1), size=size,
                             bold=True, color=self.t["text_main"], font=self.t["font_head"])

    def subtitle(self, slide, text, x=Inches(0.6), y=Inches(1.55), w=Inches(11.5), size=16):
        return self.textbox(slide, text, x, y, w, Inches(0.6), size=size, color=self.t["text_body"])

    def divider(self, slide, y, x=Inches(0.6), w=Inches(12.1)):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x + w, y)
        line.line.color.rgb = self.t["card_line"]
        line.line.width = Pt(1)
        return line

    def _header_deco(self, slide):
        h = self.st.get("header", "plain")
        W = self.prs.slide_width
        def rect(x, y, w, hh, fill):
            sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x), int(y), int(w), int(hh))
            sp.fill.solid(); sp.fill.fore_color.rgb = fill; sp.line.fill.background(); sp.shadow.inherit = False
        if h == "band":
            rect(0, 0, W, Inches(1.5), self.t["card"])
            rect(0, Inches(1.5), W, Inches(0.05), self.t["accent"])
        elif h == "tick":
            rect(Inches(0.3), Inches(0.52), Inches(0.07), Inches(0.98), self.t["accent"])

    def _decor(self, slide):
        k = self.st["decor"]
        if k == "none":
            return
        W, H = self.prs.slide_width, self.prs.slide_height
        def shp(kind, x, y, w, h, fill=None, line=None, lw=1.0):
            sp = slide.shapes.add_shape(kind, x, y, w, h)
            if fill is None: sp.fill.background()
            else:
                sp.fill.solid(); sp.fill.fore_color.rgb = fill
            if line is None: sp.line.fill.background()
            else:
                sp.line.color.rgb = line; sp.line.width = Pt(lw)
            sp.shadow.inherit = False
            return sp
        acc, card, ln = self.t["accent"], self.t["card"], self.t["card_line"]
        if k == "left_bar":
            shp(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H, fill=acc)
        elif k == "top_bar":
            shp(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.14), fill=acc)
        elif k == "corner_circle":
            d = Inches(5.2)
            shp(MSO_SHAPE.OVAL, W - Inches(2.6), H - Inches(2.6), d, d, fill=card)
        elif k == "blocks":
            shp(MSO_SHAPE.RECTANGLE, W - Inches(1.45), H - Inches(0.45), Inches(0.7), Inches(0.45), fill=acc)
            shp(MSO_SHAPE.RECTANGLE, W - Inches(0.75), H - Inches(0.45), Inches(0.75), Inches(0.45), fill=self.t["accent2"])
        elif k == "frame":
            m = Inches(0.2)
            shp(MSO_SHAPE.RECTANGLE, m, m, W - 2 * m, H - 2 * m, fill=None, line=ln, lw=2.0)

    def _titul_alt(self, kicker, baslik, sual, info_setirler, tarix_yer=""):
        """Alternativ titul düzülüşləri: 'center' və 'band'."""
        v = self.st["titul"]
        slide = self.new_slide(footer=False)
        W, H = self.prs.slide_width, self.prs.slide_height
        t = self.t
        if v == "center":
            self.textbox(slide, kicker.upper(), Inches(0.8), Inches(1.2), W - Inches(1.6), Inches(0.4),
                         size=14, bold=True, color=t["accent"], align=PP_ALIGN.CENTER)
            self.textbox(slide, baslik, Inches(0.8), Inches(1.9), W - Inches(1.6), Inches(1.8),
                         size=44, bold=True, color=t["text_main"], font=t["font_head"], align=PP_ALIGN.CENTER)
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int((W - Inches(1.4)) / 2), Inches(3.95), Inches(1.4), Inches(0.07))
            bar.fill.solid(); bar.fill.fore_color.rgb = t["accent"]; bar.line.fill.background(); bar.shadow.inherit = False
            self.textbox(slide, sual, Inches(1.5), Inches(4.2), W - Inches(3.0), Inches(0.9),
                         size=18, color=t["text_body"], align=PP_ALIGN.CENTER)
            self.textbox(slide, "\n".join(info_setirler), Inches(1.5), Inches(5.4), W - Inches(3.0), Inches(1.2),
                         size=14, color=t["text_body"], align=PP_ALIGN.CENTER)
            if tarix_yer:
                self.textbox(slide, tarix_yer, Inches(0.8), Inches(6.85), W - Inches(1.6), Inches(0.4),
                             size=12, color=t["text_dim"], align=PP_ALIGN.CENTER)
        else:  # band
            pw = Inches(5.6)
            panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, pw, H)
            panel.fill.solid(); panel.fill.fore_color.rgb = t["accent"]; panel.line.fill.background(); panel.shadow.inherit = False
            a = t["accent"]
            lum = 0.299 * a[0] + 0.587 * a[1] + 0.114 * a[2]
            on_acc = RGBColor(0x14, 0x14, 0x14) if lum > 150 else RGBColor(0xFF, 0xFF, 0xFF)
            self.textbox(slide, baslik, Inches(0.6), Inches(2.0), pw - Inches(1.0), Inches(3.4),
                         size=36, bold=True, color=on_acc, font=t["font_head"])
            self.textbox(slide, kicker.upper(), Inches(6.3), Inches(1.4), Inches(6.5), Inches(0.4),
                         size=14, bold=True, color=t["accent"])
            self.textbox(slide, sual, Inches(6.3), Inches(2.0), Inches(6.4), Inches(1.8),
                         size=22, color=t["text_main"])
            self.textbox(slide, "\n".join(info_setirler), Inches(6.3), Inches(4.4), Inches(6.4), Inches(1.6),
                         size=14, color=t["text_body"])
            if tarix_yer:
                self.textbox(slide, tarix_yer, Inches(6.3), Inches(6.7), Inches(6.4), Inches(0.4),
                             size=12, color=t["text_dim"])
        return slide

    def _footer(self, slide):
        txt = f"{self.movzu_adi}  •  {self.muellif}" if self.muellif else self.movzu_adi
        self.textbox(slide, txt, Inches(0.6), Inches(7.1), Inches(7), Inches(0.3),
                     size=10, color=self.t["text_dim"])

    def _page_number(self, slide):
        self.textbox(slide, str(self.slide_no), Inches(12.6), Inches(0.4),
                     Inches(0.5), Inches(0.4), size=12, color=self.t["text_dim"],
                     align=PP_ALIGN.RIGHT)

    def card(self, slide, x, y, w, h, fill=None, radius=None):
        radius = self.st["radius"] if radius is None else radius
        fill = fill or self.t["card"]
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = self.t["card_line"]
        shape.line.width = Pt(0.75)
        shape.shadow.inherit = False
        try:
            shape.adjustments[0] = radius
        except Exception:
            pass
        return shape

    def circle_dot(self, slide, cx, cy, d=Inches(0.22), color=None):
        color = color or self.t["accent"]
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - d // 2, cy - d // 2, d, d)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    # ---------- SLAYD ŞABLONLARI ----------

    def titul_slaydi(self, kicker, baslik, sual, info_setirler, tarix_yer="", visual_path=None):
        if self.st.get("titul", "left") != "left" and not visual_path:
            return self._titul_alt(kicker, baslik, sual, info_setirler, tarix_yer)
        slide = self.new_slide(footer=False)
        self.label(slide, kicker)

        tw = Inches(7.2) if visual_path else Inches(11)
        self.textbox(slide, baslik, Inches(0.6), Inches(1.6), tw, Inches(1.6),
                     size=42, bold=True, color=self.t["text_main"], font=self.t["font_head"])
        self.textbox(slide, sual, Inches(0.6), Inches(3.1), Inches(7.0) if visual_path else Inches(9), Inches(1),
                     size=16, color=self.t["text_body"])
        self.textbox(slide, "\n".join(info_setirler), Inches(0.6), Inches(4.3),
                     Inches(7.0) if visual_path else Inches(8), Inches(2),
                     size=13, color=self.t["text_body"])

        # ---- Sağ kənarda riyazi vizual ----
        if visual_path:
            from PIL import Image
            card_x, card_y = Inches(8.1), Inches(1.3)
            card_w, card_h = Inches(4.6), Inches(4.7)
            self.card(slide, card_x, card_y, card_w, card_h)

            with Image.open(visual_path) as im:
                iw, ih = im.size
            ratio = ih / iw
            pic_w = Inches(4.2)
            pic_h = int(pic_w * ratio)
            max_h = Inches(4.3)
            if pic_h > max_h:
                pic_h = max_h
                pic_w = int(pic_h / ratio)
            px = int(card_x + (card_w - pic_w) / 2)
            py = int(card_y + (card_h - pic_h) / 2)
            pic = slide.shapes.add_picture(visual_path, px, py, width=pic_w, height=pic_h)
            pic.shadow.inherit = False

        if tarix_yer:
            self.textbox(slide, tarix_yer, Inches(0.6), Inches(7.0), Inches(4), Inches(0.4),
                         size=12, color=self.t["text_dim"])
        return slide

    def plan_slaydi(self, baslik, maddeler):
        slide = self.new_slide()
        self.label(slide, "PLAN")
        self.title(slide, baslik)

        n = len(maddeler)
        y = Inches(2.1)
        # Sətir sayı artdıqca hündürlük avtomatik azalır (maks. 0.95")
        row_h = min(Inches(0.95), int(Inches(4.7) / n))

        for nomre, qisa, izah in maddeler:
            self.textbox(slide, nomre, Inches(0.6), y, Inches(0.6), Inches(0.5),
                         size=16, bold=True, color=self.t["accent"], font=self.t["font_head"])
            self.textbox(slide, qisa, Inches(1.3), y, Inches(2.8), Inches(0.5),
                         size=18, bold=True, color=self.t["text_main"])
            self.textbox(slide, izah, Inches(4.3), y, Inches(7.5), Inches(0.6),
                         size=15, color=self.t["text_body"])
            self.divider(slide, y + row_h - Inches(0.15))
            y += row_h
        return slide

    def movzu_slaydi(self, kicker, baslik, meqamlar):
        """Sadə bullet-list slaydı"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        y = Inches(2.1)
        for meqam in meqamlar:
            self.circle_dot(slide, Inches(0.85), y + Inches(0.15), Inches(0.12))
            self.textbox(slide, meqam, Inches(1.1), y, Inches(11), Inches(0.7),
                         size=20, color=self.t["text_body"])
            y += Inches(0.75)
        return slide

    def stat_slaydi(self, kicker, baslik, statistikalar):
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        n = len(statistikalar)
        w = Inches(11.5) // n
        x = Inches(0.9)
        for rəqəm, izah in statistikalar:
            self.textbox(slide, rəqəm, x, Inches(2.8), w, Inches(1.3),
                         size=56, bold=True, color=self.t["accent"],
                         align=PP_ALIGN.CENTER, font=self.t["font_head"])
            self.textbox(slide, izah, x, Inches(4.1), w, Inches(0.8),
                         size=16, color=self.t["text_body"], align=PP_ALIGN.CENTER)
            x += w
        return slide

    def grid_slaydi(self, kicker, baslik, kartlar, cols=3):
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        card_w = Inches(11.6) // cols - Inches(0.2)
        card_h = Inches(2.0)
        for i, (ikon, kbaslik, kizah) in enumerate(kartlar):
            row, col = divmod(i, cols)
            x = Inches(0.6) + col * (card_w + Inches(0.25))
            y = Inches(2.0) + row * (card_h + Inches(0.25))
            self.card(slide, x, y, card_w, card_h)
            self.textbox(slide, ikon, x + Inches(0.25), y + Inches(0.2), Inches(0.8), Inches(0.5),
                         size=22, bold=True, color=self.t["accent"])
            self.textbox(slide, kbaslik, x + Inches(0.25), y + Inches(0.75), card_w - Inches(0.5), Inches(0.4),
                         size=14, bold=True, color=self.t["text_main"])
            self.textbox(slide, kizah, x + Inches(0.25), y + Inches(1.1), card_w - Inches(0.5), Inches(0.85),
                         size=13, color=self.t["text_body"])
        return slide

    def zigzag_slaydi(self, kicker, baslik, bloklar):
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        y = Inches(2.0)
        for abaslik, ametn, tərəf in bloklar:
            dot_x = Inches(0.9) if tərəf == "left" else Inches(10.5)
            text_x = Inches(1.4) if tərəf == "left" else Inches(1.0)
            self.circle_dot(slide, dot_x, y + Inches(0.2), Inches(0.18))
            self.textbox(slide, abaslik, text_x, y, Inches(10), Inches(0.4),
                         size=16, bold=True, color=self.t["text_main"])
            self.textbox(slide, ametn, text_x, y + Inches(0.45), Inches(10), Inches(0.5),
                         size=13, color=self.t["text_body"])
            y += Inches(1.1)
        return slide

    def proses_axini_slaydi(self, kicker, baslik, addimlar, neticə_formula=""):
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        n = len(addimlar)
        card_w = Inches(2.75)
        gap = Inches(0.3)
        total = n * card_w + (n - 1) * gap
        start_x = (self.prs.slide_width - total) // 2
        y = Inches(2.0)
        card_h = Inches(3.0)
        for i, (nomre, abaslik, aizah) in enumerate(addimlar):
            x = start_x + i * (card_w + gap)
            self.card(slide, x, y, card_w, card_h)
            self.textbox(slide, nomre, x + Inches(0.25), y + Inches(0.2),
                         Inches(1), Inches(0.4), size=16, bold=True,
                         color=self.t["accent"], font=self.t["font_head"])
            self.textbox(slide, abaslik, x + Inches(0.25), y + Inches(0.75),
                         card_w - Inches(0.5), Inches(0.8), size=16, bold=True, color=self.t["text_main"])
            self.textbox(slide, aizah, x + Inches(0.25), y + Inches(1.5),
                         card_w - Inches(0.5), Inches(1.4), size=13, color=self.t["text_body"])
            if i < n - 1:
                ax = x + card_w + gap // 2
                self.textbox(slide, "→", ax - Inches(0.2), y + card_h // 2 - Inches(0.25),
                             Inches(0.4), Inches(0.5), size=22, bold=True,
                             color=self.t["text_dim"], align=PP_ALIGN.CENTER)
        if neticə_formula:
            self.card(slide, start_x, y + card_h + Inches(0.4), total, Inches(0.9), fill=self.t["card"])
            self.textbox(slide, "NƏTİCƏ", start_x + Inches(0.3),
                         y + card_h + Inches(0.65), Inches(1.5), Inches(0.4),
                         size=12, bold=True, color=self.t["accent"])
            self.textbox(slide, neticə_formula, start_x + Inches(1.8),
                         y + card_h + Inches(0.55), total - Inches(2), Inches(0.6),
                         size=20, bold=True, color=self.t["text_main"], font=self.t["font_head"])
        return slide

    def timeline_slaydi(self, kicker, baslik, hadiseler, esas_fikir=""):
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        n = len(hadiseler)
        line_y = Inches(3.5)
        self.divider(slide, line_y, x=Inches(1), w=Inches(11.3))
        gap = Inches(11.3) // n
        for i, (tarix, hbaslik, hizah) in enumerate(hadiseler):
            cx = Inches(1) + gap * i + gap // 2
            self.circle_dot(slide, cx, line_y, Inches(0.22))
            self.textbox(slide, tarix, cx - Inches(0.9), line_y - Inches(0.65),
                         Inches(1.8), Inches(0.4), size=14, bold=True,
                         color=self.t["accent"], align=PP_ALIGN.CENTER)
            self.textbox(slide, hbaslik, cx - Inches(1), line_y + Inches(0.35),
                         Inches(2), Inches(0.4), size=14, bold=True,
                         color=self.t["text_main"], align=PP_ALIGN.CENTER)
            self.textbox(slide, hizah, cx - Inches(1), line_y + Inches(0.75),
                         Inches(2), Inches(1.0), size=13, color=self.t["text_body"],
                         align=PP_ALIGN.CENTER)
        if esas_fikir:
            self.card(slide, Inches(0.6), Inches(5.7), Inches(12.1), Inches(1))
            self.textbox(slide, "Əsas fikir", Inches(1), Inches(5.95), Inches(1.5),
                         Inches(0.4), size=13, bold=True, color=self.t["accent"])
            self.textbox(slide, esas_fikir, Inches(2.6), Inches(5.95), Inches(9.8),
                         Inches(0.7), size=14, color=self.t["text_main"])
        return slide

    def formula_slaydi(self, kicker, baslik, formula, aciqlamalar):
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        self.card(slide, Inches(0.6), Inches(1.8), Inches(12.1), Inches(1.6))
        self.textbox(slide, formula, Inches(1), Inches(2.1), Inches(11.3), Inches(1),
                     size=30, bold=True, color=self.t["text_main"], font=self.t["font_head"],
                     align=PP_ALIGN.CENTER)
        n = len(aciqlamalar)
        card_w = (Inches(12.1) - Inches(0.3) * (n - 1)) // n
        x = Inches(0.6)
        for hissə, izah in aciqlamalar:
            self.card(slide, x, Inches(3.8), card_w, Inches(1.4))
            self.textbox(slide, hissə, x, Inches(4.0), card_w, Inches(0.5),
                         size=16, bold=True, color=self.t["accent"], align=PP_ALIGN.CENTER,
                         font=self.t["font_head"])
            self.textbox(slide, izah, x + Inches(0.15), Inches(4.55), card_w - Inches(0.3),
                         Inches(0.7), size=14, color=self.t["text_body"], align=PP_ALIGN.CENTER)
            x += card_w + Inches(0.3)
        return slide

    def cedvel_slaydi(self, kicker, baslik, basliqlar, setirler):
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        rows, cols = len(setirler) + 1, len(basliqlar)
        table_shape = slide.shapes.add_table(rows, cols, Inches(0.6), Inches(2.0),
                                              Inches(12.1), Inches(0.6 * rows))
        table = table_shape.table
        for i, h in enumerate(basliqlar):
            cell = table.cell(0, i)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = self.t["accent"]
            p = cell.text_frame.paragraphs[0]
            p.font.bold = True
            p.font.color.rgb = self.t["bg"]
            p.font.size = Pt(16)
        for r, row_data in enumerate(setirler, start=1):
            for c, val in enumerate(row_data):
                cell = table.cell(r, c)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.t["card"] if r % 2 else self.t["bg"]
                cell.text_frame.paragraphs[0].font.color.rgb = self.t["text_body"]
                cell.text_frame.paragraphs[0].font.size = Pt(14)
        return slide

    def sitat_slaydi(self, sitat, muellif=""):
        slide = self.new_slide(footer=False)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.t["accent2"]
        big_quote = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(2), Inches(2))
        big_quote.text_frame.text = '"'
        big_quote.text_frame.paragraphs[0].font.size = Pt(120)
        big_quote.text_frame.paragraphs[0].font.bold = True
        big_quote.text_frame.paragraphs[0].font.color.rgb = self.t["accent"]
        self.textbox(slide, sitat, Inches(1.2), Inches(2.6), Inches(10.9), Inches(2),
                     size=30, bold=True, color=self.t["text_main"])
        if muellif:
            self.textbox(slide, f"— {muellif}", Inches(1.2), Inches(5.2), Inches(10), Inches(0.5),
                         size=16, color=self.t["accent"])
        return slide

    def teşekkur_slaydi(self, formula=""):
        slide = self.new_slide(footer=False)
        self.divider(slide, Inches(1.0))
        self.textbox(slide, "TƏŞƏKKÜR EDİRİK", Inches(0.6), Inches(2.8), Inches(12.1),
                     Inches(1), size=38, bold=True, color=self.t["text_main"],
                     align=PP_ALIGN.CENTER, font=self.t["font_head"])
        self.textbox(slide, "Diqqətinizə görə təşəkkür edirik.", Inches(0.6), Inches(3.7),
                     Inches(12.1), Inches(0.5), size=16, color=self.t["text_body"],
                     align=PP_ALIGN.CENTER)
        if formula:
            card_w = Inches(6)
            x = (self.prs.slide_width - card_w) // 2
            self.card(slide, x, Inches(4.4), card_w, Inches(0.9))
            self.textbox(slide, formula, x, Inches(4.6), card_w, Inches(0.5),
                         size=18, bold=True, color=self.t["accent"], align=PP_ALIGN.CENTER,
                         font=self.t["font_head"])
        self._footer(slide)
        self._page_number(slide)
        return slide

    def bos_slayd(self):
        """Tam boş slayd — özünüz manual şəkildə doldurmaq istəsəniz"""
        return self.new_slide()

    # ---------- BUILD-REVEAL (ardıcıl açılma effekti) ----------

    def build_reveal(self, render_func, items, **kwargs):
        """
        render_func: lambda subset: self.some_slaydi(..., subset, ...)
        Hər addımda 1 element artıraraq N slayd yaradır.
        """
        for i in range(1, len(items) + 1):
            render_func(items[:i], **kwargs)

    # ---------- YEKUN ----------

    # ============================================================
    # YENİ ƏLAVƏLƏR — Riyazi Analiz Mövzusu Üçün
    # ============================================================

    def iki_kart_slaydi(self, kicker, baslik, sol_baslik, sol_meqamlar,
                         sag_baslik, sag_meqamlar, izah=""):
        """İki bərabər kart - müqayisə və ya paralel məlumat üçün"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        if izah:
            self.subtitle(slide, izah)

        y0 = Inches(2.2) if izah else Inches(1.8)
        w, h = Inches(5.9), Inches(4.3)
        data = [(sol_baslik, sol_meqamlar, Inches(0.6)), (sag_baslik, sag_meqamlar, Inches(6.85))]
        for kbaslik, meqamlar, x0 in data:
            self.card(slide, x0, y0, w, h)
            self.textbox(slide, kbaslik, x0 + Inches(0.3), y0 + Inches(0.25), w - Inches(0.6), Inches(0.5),
                         size=17, bold=True, color=self.t["accent"])
            yy = y0 + Inches(0.95)
            for m in meqamlar:
                self.circle_dot(slide, x0 + Inches(0.45), yy + Inches(0.12), Inches(0.1))
                self.textbox(slide, m, x0 + Inches(0.65), yy, w - Inches(1), Inches(0.65),
                             size=15, color=self.t["text_body"])
                yy += Inches(0.7)
        return slide

    def iddia_slaydi(self, kicker, baslik, iddia, izah="", simvol=""):
        """Böyük, mərkəzləşdirilmiş riyazi iddia/teorem vurğusu"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        self.card(slide, Inches(1.3), Inches(2.3), Inches(10.7), Inches(2.6))
        yy = Inches(2.55)
        if simvol:
            self.textbox(slide, simvol, Inches(1.6), yy, Inches(10.1), Inches(0.7),
                         size=22, bold=True, color=self.t["accent"],
                         align=PP_ALIGN.CENTER, font=self.t["font_head"])
            yy = Inches(3.35)
        self.textbox(slide, iddia, Inches(1.6), yy, Inches(10.1), Inches(0.8),
                     size=18, bold=True, color=self.t["text_main"], align=PP_ALIGN.CENTER)
        if izah:
            self.textbox(slide, izah, Inches(1.8), yy + Inches(0.9), Inches(9.7), Inches(0.8),
                         size=13, color=self.t["text_body"], align=PP_ALIGN.CENTER)
        return slide

    def nomreli_numune_slaydi(self, kicker, baslik, verilenler, addimlar, cavab):
        """Addım-addım ədədi həll nümunəsi"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)

        self.card(slide, Inches(0.6), Inches(1.8), Inches(12.1), Inches(0.9))
        self.textbox(slide, "Verilir:", Inches(1), Inches(2.05), Inches(1.3),
                     Inches(0.4), size=12, bold=True, color=self.t["text_body"])
        self.textbox(slide, verilenler, Inches(2.4), Inches(1.95), Inches(9.5),
                     Inches(0.6), size=20, bold=True, color=self.t["text_main"],
                     font=self.t["font_head"])

        n = len(addimlar)
        card_w = (Inches(12.1) - Inches(0.3) * (n - 1)) // n
        x = Inches(0.6)
        for nomre, adi, hesab in addimlar:
            self.card(slide, x, Inches(3.1), card_w, Inches(1.4))
            self.textbox(slide, nomre, x + Inches(0.2), Inches(3.25), Inches(0.6),
                         Inches(0.4), size=13, bold=True, color=self.t["accent"])
            self.textbox(slide, adi, x + Inches(0.7), Inches(3.25), card_w - Inches(0.9),
                         Inches(0.4), size=14, bold=True, color=self.t["text_main"])
            self.textbox(slide, hesab, x, Inches(3.8), card_w, Inches(0.6),
                         size=15, color=self.t["text_body"], align=PP_ALIGN.CENTER,
                         font=self.t["font_body"])
            x += card_w + Inches(0.3)

        self.card(slide, Inches(0.6), Inches(4.9), Inches(12.1), Inches(1))
        self.textbox(slide, "CAVAB", Inches(1), Inches(5.2), Inches(1.3), Inches(0.5),
                     size=12, bold=True, color=self.t["accent"])
        self.textbox(slide, cavab, Inches(2.4), Inches(5.1), Inches(9.5), Inches(0.7),
                     size=22, bold=True, color=self.t["text_main"], font=self.t["font_head"])
        return slide

    def diagram_slaydi(self, kicker, baslik, izah, diagram_func, caption=""):
        """Tam enli diaqram slaydı - istənilən çəkmə funksiyasını qəbul edir"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        if izah:
            self.subtitle(slide, izah)
        y0 = Inches(2.3) if izah else Inches(1.9)
        card_h = Inches(4.1)
        self.card(slide, Inches(0.8), y0, Inches(11.7), card_h)
        diagram_func(slide, Inches(1.3), y0 + Inches(0.3), Inches(10.7), card_h - Inches(0.6))
        if caption:
            self.textbox(slide, caption, Inches(0.8), y0 + card_h + Inches(0.15), Inches(11.7), Inches(0.4),
                         size=13, bold=True, color=self.t["accent"], align=PP_ALIGN.CENTER)
        return slide

    # ---------- XAM DİAQRAM ÇƏKƏNLƏR (diagram_slaydi ilə istifadə üçün) ----------

    def limit_convergence_diagram(self, slide, x, y, w, h):
        """Ədəd xəttində ardıcıllığın limitə necə yığıldığını göstərir"""
        line_y = int(y + h * 0.65)
        main_line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x), line_y, int(x + w), line_y)
        main_line.line.color.rgb = self.t["text_dim"]
        main_line.line.width = Pt(1.5)

        limit_x = int(x + w * 0.88)
        band_w = Inches(0.45)
        self.card(slide, limit_x - band_w // 2, int(y + h * 0.15), band_w, int(h * 0.6),
                  fill=self.t["accent2"], radius=0.15)

        lim_line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                               limit_x, int(y + h * 0.05),
                                               limit_x, int(y + h * 0.95))
        lim_line.line.color.rgb = self.t["accent"]
        lim_line.line.width = Pt(2)

        self.textbox(slide, "a", limit_x - Inches(0.2), int(y), Inches(0.4), Inches(0.3),
                     size=14, bold=True, color=self.t["accent"], align=PP_ALIGN.CENTER)

        n_points = 5
        start_x = x + w * 0.06
        for i in range(n_points):
            frac = i / (n_points - 1)
            px = int(start_x + (limit_x - start_x) * (frac ** 1.8))
            self.circle_dot(slide, px, line_y, Inches(0.12))
            self.textbox(slide, f"x{i+1}", px - Inches(0.3), line_y + Inches(0.15), Inches(0.6), Inches(0.3),
                         size=10, color=self.t["text_body"], align=PP_ALIGN.CENTER)

    def squeeze_diagram(self, slide, x, y, w, h):
        """Sıxışdırma (iki polis) teoremini göstərən sxem"""
        end_x = int(x + w * 0.88)
        end_y = int(y + h * 0.5)
        start_x = int(x + w * 0.05)

        def draw_path(points, color, width=2.5):
            for i in range(len(points) - 1):
                conn = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT,
                    int(points[i][0]), int(points[i][1]),
                    int(points[i+1][0]), int(points[i+1][1])
                )
                conn.line.color.rgb = color
                conn.line.width = Pt(width)

        mid_x = int(x + w * 0.45)
        upper = [(start_x, y + h*0.1), (mid_x, y + h*0.22), (end_x, end_y)]
        lower = [(start_x, y + h*0.9), (mid_x, y + h*0.78), (end_x, end_y)]
        middle = [(start_x, y + h*0.5), (mid_x, y + h*0.5), (end_x, end_y)]

        draw_path(upper, self.t["accent2"])
        draw_path(lower, self.t["accent2"])
        draw_path(middle, self.t["accent"], width=3.5)

        self.circle_dot(slide, end_x, end_y, Inches(0.15))
        self.textbox(slide, "a", end_x + Inches(0.15), end_y - Inches(0.15), Inches(0.4), Inches(0.3),
                     size=14, bold=True, color=self.t["accent"])

        self.textbox(slide, "yₙ", int(start_x - Inches(0.05)), int(y + h*0.02), Inches(0.5), Inches(0.3),
                     size=12, bold=True, color=self.t["accent2"])
        self.textbox(slide, "zₙ", int(start_x - Inches(0.05)), int(y + h*0.42), Inches(0.5), Inches(0.3),
                     size=12, bold=True, color=self.t["accent"])
        self.textbox(slide, "xₙ", int(start_x - Inches(0.05)), int(y + h*0.82), Inches(0.5), Inches(0.3),
                     size=12, bold=True, color=self.t["accent2"])

    def rising_plateau_diagram(self, slide, x, y, w, h, label_y="Konsentrasiya", label_x="Zaman"):
        """Artıb müəyyən həddə sabitləşən əyri (dərman, bank balansı və s. üçün)"""
        base_y = int(y + h * 0.85)
        top_y = int(y + h * 0.15)

        y_axis = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x), top_y, int(x), base_y)
        y_axis.line.color.rgb = self.t["text_dim"]
        y_axis.line.width = Pt(1.25)

        x_axis = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x), base_y, int(x + w), base_y)
        x_axis.line.color.rgb = self.t["text_dim"]
        x_axis.line.width = Pt(1.25)

        plateau_y = int(y + h * 0.3)
        plateau = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                              int(x + Inches(0.1)), plateau_y,
                                              int(x + w), plateau_y)
        plateau.line.color.rgb = self.t["accent2"]
        plateau.line.width = Pt(1)

        n_seg = 8
        pts = []
        for i in range(n_seg + 1):
            frac = i / n_seg
            px = x + Inches(0.1) + frac * (w - Inches(0.2))
            val = 1 - math.exp(-3.2 * frac)
            py = base_y - val * (base_y - plateau_y)
            pts.append((int(px), int(py)))

        for i in range(len(pts) - 1):
            seg = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                              pts[i][0], pts[i][1], pts[i+1][0], pts[i+1][1])
            seg.line.color.rgb = self.t["accent"]
            seg.line.width = Pt(3)

        self.textbox(slide, label_x, int(x + w - Inches(1)), base_y + Inches(0.1), Inches(1.2), Inches(0.3),
                     size=11, color=self.t["text_dim"])
        self.textbox(slide, label_y, int(x - Inches(0.1)), top_y - Inches(0.4), Inches(2.8), Inches(0.3),
                     size=11, color=self.t["text_dim"])

    def picture_slaydi(self, kicker, baslik, izah, image_path, caption="", img_width=Inches(9.5)):
        """Tam en, mərkəzləşdirilmiş şəkil/qrafik slaydı (matplotlib PNG üçün)"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        if izah:
            self.subtitle(slide, izah)
        y0 = Inches(2.3) if izah else Inches(1.9)

        with Image.open(image_path) as im:
            iw, ih = im.size
        ratio = ih / iw
        w = img_width
        h = int(w * ratio)
        max_h = Inches(4.3)
        if h > max_h:
            h = max_h
            w = int(h / ratio)

        x = (self.prs.slide_width - w) // 2
        pic = slide.shapes.add_picture(image_path, x, y0, width=w, height=h)
        pic.shadow.inherit = False

        if caption:
            self.textbox(slide, caption, Inches(0.8), y0 + h + Inches(0.15), Inches(11.7), Inches(0.4),
                         size=13, bold=True, color=self.t["accent"], align=PP_ALIGN.CENTER)
        return slide

    def icon_grid_slaydi(self, kicker, baslik, items, cols=3, placeholders=None):
        """items: [(icon_path, başlıq, izah), ...]
        placeholders: [(etiket, hint), ...] — kartların altında şəkil yeri (istəyə bağlı)"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)

        n = len(items)
        card_w = Inches(11.6) // cols - Inches(0.2)
        card_h = Inches(2.1)
        rows = (n + cols - 1) // cols

        for i, (icon_path, kbaslik, kizah) in enumerate(items):
            row, col = divmod(i, cols)
            x = Inches(0.6) + col * (card_w + Inches(0.25))
            y = Inches(2.0) + row * (card_h + Inches(0.25))

            self.card(slide, x, y, card_w, card_h)
            icon_size = Inches(0.7)
            try:
                pic = slide.shapes.add_picture(icon_path, x + Inches(0.25), y + Inches(0.22),
                                                width=icon_size, height=icon_size)
                pic.shadow.inherit = False
            except Exception:
                pass
            self.textbox(slide, kbaslik, x + Inches(1.1), y + Inches(0.22), card_w - Inches(1.3), Inches(0.5),
                         size=14, bold=True, color=self.t["text_main"])
            self.textbox(slide, kizah, x + Inches(0.25), y + Inches(1.05), card_w - Inches(0.5), Inches(0.95),
                         size=11, color=self.t["text_body"])

        # ---- Kartların altında şəkil yerləri ----
        if placeholders:
            ph_h = Inches(1.75)
            ph_y = int(Inches(2.0) + rows * (card_h + Inches(0.25)) + Inches(0.1))
            for j, (etiket, hint) in enumerate(placeholders[:cols]):
                x = Inches(0.6) + j * (card_w + Inches(0.25))
                box = self.card(slide, x, ph_y, card_w, ph_h, fill=self.t["bg"])
                box.line.color.rgb = self.t["accent"]
                box.line.width = Pt(1.25)

                self.textbox(slide, "🖼", x, int(ph_y + Inches(0.2)), card_w, Inches(0.5),
                             size=24, color=self.t["accent"], align=PP_ALIGN.CENTER)
                self.textbox(slide, hint, x + Inches(0.2), int(ph_y + Inches(0.8)),
                             card_w - Inches(0.4), Inches(0.8),
                             size=9, color=self.t["text_dim"], align=PP_ALIGN.CENTER)
                self.textbox(slide, etiket, x, int(ph_y + ph_h + Inches(0.05)), card_w, Inches(0.35),
                             size=13, bold=True, color=self.t["text_main"], align=PP_ALIGN.CENTER)
        return slide

    def section_slaydi(self, kicker, baslik, izah="", icon_path=None, addim_etiketi=""):
        """Bölmə keçidi — təkmilləşdirilmiş dizayn (dekorativ dairə + ikon)"""
        slide = self.new_slide(footer=False)

        # Böyük dekorativ dairə (arxa fonda, sağ üst)
        circle_d = Inches(7.5)
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            self.prs.slide_width - Inches(4.5), Inches(-2.2),
            circle_d, circle_d
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = self.t["card"]
        circle.line.fill.background()
        circle.shadow.inherit = False

        # Kiçik vurğu dairəsi
        small_circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            self.prs.slide_width - Inches(2.3), Inches(4.6),
            Inches(0.9), Inches(0.9)
        )
        small_circle.fill.solid()
        small_circle.fill.fore_color.rgb = self.t["accent"]
        small_circle.line.fill.background()
        small_circle.shadow.inherit = False

        if addim_etiketi:
            self.textbox(slide, addim_etiketi, Inches(0.6), Inches(0.5), Inches(3), Inches(0.5),
                         size=14, bold=True, color=self.t["accent"])

        self.label(slide, kicker, x=Inches(0.6), y=Inches(2.5))
        self.textbox(slide, baslik, Inches(0.6), Inches(3.0), Inches(7.5), Inches(1.3),
                     size=34, bold=True, color=self.t["text_main"], font=self.t["font_head"])
        if izah:
            self.textbox(slide, izah, Inches(0.6), Inches(4.2), Inches(7), Inches(0.7),
                         size=15, color=self.t["text_body"])
        self.divider(slide, Inches(2.9), x=Inches(0.6), w=Inches(2.5))

        if icon_path:
            try:
                pic = slide.shapes.add_picture(icon_path, self.prs.slide_width - Inches(4.9),
                                                Inches(1.3), height=Inches(2.6))
                pic.shadow.inherit = False
            except Exception:
                pass

        self._footer(slide)
        self._page_number(slide)
        return slide
    
    def _contrast_color(self, rgb):
        """Fon rənginə görə oxunaqlı mətn rəngi (ağ/qara) seçir"""
        r, g, b = rgb[0], rgb[1], rgb[2]
        luminance = (0.299*r + 0.587*g + 0.114*b) / 255
        return RGBColor(0x00, 0x00, 0x00) if luminance > 0.6 else RGBColor(0xFF, 0xFF, 0xFF)

    def hook_slaydi(self, kicker, sual, alt_metn=""):
        """Tam ekran, rəngli fonlu açılış sualı — diqqət cəlb etmək üçün"""
        slide = self.new_slide(footer=False)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.t["accent2"]
        txt_color = self._contrast_color(self.t["accent2"])

        self.textbox(slide, kicker.upper(), Inches(0.8), Inches(0.6), Inches(8), Inches(0.4),
                     size=13, bold=True, color=txt_color)
        self.textbox(slide, sual, Inches(0.8), Inches(2.5), Inches(11.7), Inches(2.3),
                     size=34, bold=True, color=txt_color, font=self.t["font_head"])
        if alt_metn:
            self.textbox(slide, alt_metn, Inches(0.8), Inches(4.9), Inches(10.5), Inches(0.8),
                         size=16, color=txt_color)
        return slide

    def image_placeholder_slaydi(self, kicker, baslik, izah, hint_text, caption="", img_ratio=16/9):
        """Tək, böyük şəkil yeri — real foto üçün"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        if izah:
            self.subtitle(slide, izah)
        y0 = Inches(2.3) if izah else Inches(1.9)
        w = Inches(9.5)
        h = int(w / img_ratio)
        max_h = Inches(4.0)
        if h > max_h:
            h = max_h
            w = int(h * img_ratio)
        x = (self.prs.slide_width - w) // 2

        box = self.card(slide, x, y0, w, h, fill=self.t["bg"])
        box.line.color.rgb = self.t["accent"]
        box.line.width = Pt(1.5)

        self.textbox(slide, "🖼", x, int(y0 + h*0.32), w, Inches(0.7),
                     size=40, color=self.t["accent"], align=PP_ALIGN.CENTER)
        self.textbox(slide, "ŞƏKİL ƏLAVƏ EDİN", x, int(y0 + h*0.32 + Inches(0.6)), w, Inches(0.4),
                     size=14, bold=True, color=self.t["text_body"], align=PP_ALIGN.CENTER)
        self.textbox(slide, hint_text, x + Inches(0.4), int(y0 + h*0.32 + Inches(1.0)), w - Inches(0.8), Inches(0.5),
                     size=11, color=self.t["text_dim"], align=PP_ALIGN.CENTER)

        if caption:
            self.textbox(slide, caption, Inches(0.8), y0 + h + Inches(0.15), Inches(11.7), Inches(0.4),
                         size=13, bold=True, color=self.t["accent"], align=PP_ALIGN.CENTER)
        return slide

    def iki_placeholder_slaydi(self, kicker, baslik, sol_etiket, sol_hint, sag_etiket, sag_hint):
        """İki şəkil yeri yan-yana — portretlər və ya müqayisə fotoları üçün"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        y0 = Inches(2.0)
        w, h = Inches(3.5), Inches(4.2)
        data = [(sol_etiket, sol_hint, Inches(1.8)), (sag_etiket, sag_hint, Inches(8.05))]
        for etiket, hint, x in data:
            box = self.card(slide, x, y0, w, h, fill=self.t["bg"])
            box.line.color.rgb = self.t["accent"]
            box.line.width = Pt(1.5)
            self.textbox(slide, "🖼", x, int(y0 + h*0.32), w, Inches(0.7),
                         size=36, color=self.t["accent"], align=PP_ALIGN.CENTER)
            self.textbox(slide, hint, x + Inches(0.3), int(y0 + h*0.32 + Inches(0.6)), w - Inches(0.6), Inches(0.7),
                         size=10, color=self.t["text_dim"], align=PP_ALIGN.CENTER)
            self.textbox(slide, etiket, x, y0 + h + Inches(0.1), w, Inches(0.4),
                         size=14, bold=True, color=self.t["text_main"], align=PP_ALIGN.CENTER)
        return slide

    def split_media_slaydi(self, kicker, baslik, izah, placeholder_hint, chart_path,
                            left_caption="", right_caption=""):
        """Sol: real şəkil yeri | Sağ: hazır qrafik (chart) — bir slaydda ikisi"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)
        if izah:
            self.subtitle(slide, izah)
        y0 = Inches(2.2) if izah else Inches(1.8)
        w, h = Inches(5.7), Inches(4.0)

        x1 = Inches(0.7)
        box = self.card(slide, x1, y0, w, h, fill=self.t["bg"])
        box.line.color.rgb = self.t["accent"]
        box.line.width = Pt(1.5)
        self.textbox(slide, "🖼", x1, int(y0 + h*0.3), w, Inches(0.7),
                     size=34, color=self.t["accent"], align=PP_ALIGN.CENTER)
        self.textbox(slide, "ŞƏKİL ƏLAVƏ EDİN", x1, int(y0 + h*0.3 + Inches(0.55)), w, Inches(0.35),
                     size=12, bold=True, color=self.t["text_body"], align=PP_ALIGN.CENTER)
        self.textbox(slide, placeholder_hint, x1 + Inches(0.3), int(y0 + h*0.3 + Inches(0.95)), w - Inches(0.6), Inches(0.6),
                     size=10, color=self.t["text_dim"], align=PP_ALIGN.CENTER)
        if left_caption:
            self.textbox(slide, left_caption, x1, y0 + h + Inches(0.1), w, Inches(0.35),
                         size=11, bold=True, color=self.t["accent"], align=PP_ALIGN.CENTER)

        x2 = Inches(6.95)
        self.card(slide, x2, y0, w, h)
        with Image.open(chart_path) as im:
            iw, ih = im.size
        ratio = ih / iw
        pw = w - Inches(0.3)
        ph = int(pw * ratio)
        max_ph = h - Inches(0.3)
        if ph > max_ph:
            ph = max_ph
            pw = int(ph / ratio)
        px = x2 + (w - pw) // 2
        py = y0 + (h - ph) // 2
        pic = slide.shapes.add_picture(chart_path, px, py, width=pw, height=ph)
        pic.shadow.inherit = False
        if right_caption:
            self.textbox(slide, right_caption, x2, y0 + h + Inches(0.1), w, Inches(0.35),
                         size=11, bold=True, color=self.t["accent"], align=PP_ALIGN.CENTER)
        return slide

    def test_sualu_slaydi(self, kicker, sual, cavablar, dogru_index=None):
        """Sınaq sualı — auditoriyanı cəlb etmək üçün, düzgün cavab vurğulanır"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.textbox(slide, sual, Inches(0.6), Inches(0.9), Inches(11.5), Inches(1.1),
                     size=24, bold=True, color=self.t["text_main"], font=self.t["font_head"])
        herflər = ["A", "B", "C", "D", "E"]
        y = Inches(2.3)
        for i, cavab in enumerate(cavablar):
            is_correct = (dogru_index == i)
            fill = self.t["accent2"] if is_correct else self.t["card"]
            box = self.card(slide, Inches(1.2), y, Inches(10.9), Inches(0.75), fill=fill)
            txt_color = self._contrast_color(fill) if is_correct else self.t["text_main"]
            tf = box.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.text = f"{herflər[i]})  {cavab}"
            p.font.size = Pt(16)
            p.font.bold = is_correct
            p.font.color.rgb = txt_color
            y += Inches(0.95)
        return slide

    def cox_placeholder_slaydi(self, kicker, baslik, placeholders, cols=2, en_boy=None):
        """placeholders=[(etiket, hint), ...]
        en_boy: qutunun en/hündürlük nisbəti (məs. 1000/1200 → portret)"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)

        n = len(placeholders)
        rows = (n + cols - 1) // cols
        gap = Inches(0.3)
        total_w = Inches(12.1)
        w = (total_w - gap * (cols - 1)) // cols
        h = int(w / en_boy) if en_boy else Inches(1.9)
        row_gap = Inches(0.45)

        y0 = Inches(1.85)
        for i, (etiket, hint) in enumerate(placeholders):
            row, col = divmod(i, cols)
            x = Inches(0.6) + col * (w + gap)
            y = y0 + row * (h + row_gap)

            box = self.card(slide, x, y, w, h, fill=self.t["bg"])
            box.line.color.rgb = self.t["accent"]
            box.line.width = Pt(1.25)

            self.textbox(slide, "🖼", x, int(y + h * 0.3), w, Inches(0.55),
                         size=26, color=self.t["accent"], align=PP_ALIGN.CENTER)
            self.textbox(slide, hint, x + Inches(0.2), int(y + h * 0.3 + Inches(0.55)),
                         w - Inches(0.4), Inches(0.9),
                         size=9, color=self.t["text_dim"], align=PP_ALIGN.CENTER)
            self.textbox(slide, etiket, x, y + h + Inches(0.05), w, Inches(0.35),
                         size=13, bold=True, color=self.t["text_main"], align=PP_ALIGN.CENTER)
        return slide

    def proses_qrafikli_slaydi(self, kicker, baslik, addimlar, image_path, neticə_formula=""):
        """Sol tərəfdə addım kartları, sağ tərəfdə real qrafik/şəkil"""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)

        # Sağ: şəkil
        img_w = Inches(5.6)
        with Image.open(image_path) as im:
            iw, ih = im.size
        ratio = ih / iw
        img_h = int(img_w * ratio)
        max_h = Inches(4.3)
        if img_h > max_h:
            img_h = max_h
            img_w = int(img_h / ratio)

        x_img = Inches(6.9)
        y_img = Inches(1.95)
        self.card(slide, x_img - Inches(0.15), y_img - Inches(0.15),
                  img_w + Inches(0.3), img_h + Inches(0.3))
        pic = slide.shapes.add_picture(image_path, x_img, y_img, width=img_w, height=img_h)
        pic.shadow.inherit = False

        # Sol: addım kartları
        step_h = Inches(0.85)
        gap = Inches(0.15)
        y = Inches(1.95)
        for nomre, abaslik, aizah in addimlar:
            self.card(slide, Inches(0.6), y, Inches(5.8), step_h)
            self.textbox(slide, nomre, Inches(0.8), y + Inches(0.1), Inches(0.6), Inches(0.4),
                         size=14, bold=True, color=self.t["accent"], font=self.t["font_head"])
            self.textbox(slide, abaslik, Inches(1.35), y + Inches(0.06), Inches(4.9), Inches(0.4),
                         size=13, bold=True, color=self.t["text_main"])
            self.textbox(slide, aizah, Inches(1.35), y + Inches(0.44), Inches(4.9), Inches(0.4),
                         size=11, color=self.t["text_body"])
            y += step_h + gap

        if neticə_formula:
            self.card(slide, Inches(0.6), y, Inches(5.8), Inches(0.65), fill=self.t["accent2"])
            txt_color = self._contrast_color(self.t["accent2"])
            self.textbox(slide, neticə_formula, Inches(0.8), y + Inches(0.12), Inches(5.4), Inches(0.4),
                         size=14, bold=True, color=txt_color, align=PP_ALIGN.CENTER)
        return slide

    def portret_slaydi(self, kicker, baslik, shexsler):
        """shexsler: [(ad, iller, izah, hint), ...]
        2 nəfər → solda portret, sağda mətn
        3-4 nəfər → portretlər yan-yana, mətn altda
        Şəkil yeri həmişə 1000×1200 (5:6) nisbətindədir."""
        slide = self.new_slide()
        self.label(slide, kicker)
        self.title(slide, baslik)

        n = len(shexsler)
        total_w = Inches(12.1)
        gap = Inches(0.3)

        def sekil_yeri(x, y, w, h, hint):
            box = self.card(slide, x, y, w, h, fill=self.t["bg"])
            box.line.color.rgb = self.t["accent"]
            box.line.width = Pt(1.5)
            self.textbox(slide, "🖼", x, int(y + h * 0.28), w, Inches(0.6),
                         size=28, color=self.t["accent"], align=PP_ALIGN.CENTER)
            self.textbox(slide, hint + "\n(1000×1200px)", x + Inches(0.15),
                         int(y + h * 0.28 + Inches(0.65)), w - Inches(0.3), Inches(0.8),
                         size=9, color=self.t["text_dim"], align=PP_ALIGN.CENTER)

        if n <= 2:
            half_w = (total_w - gap * (n - 1)) // n
            pw = Inches(2.9)
            ph = int(pw * 1.2)
            y0 = Inches(2.0)
            for i, (ad, iller, izah, hint) in enumerate(shexsler):
                x = Inches(0.6) + i * (half_w + gap)
                sekil_yeri(x, y0, pw, ph, hint)
                tx = x + pw + Inches(0.25)
                tw = half_w - pw - Inches(0.25)
                self.textbox(slide, ad, tx, y0 + Inches(0.1), tw, Inches(0.9),
                             size=18, bold=True, color=self.t["text_main"],
                             font=self.t["font_head"])
                self.textbox(slide, iller, tx, y0 + Inches(1.05), tw, Inches(0.4),
                             size=12, bold=True, color=self.t["accent"])
                self.textbox(slide, izah, tx, y0 + Inches(1.55), tw, Inches(1.9),
                             size=12, color=self.t["text_body"])
        else:
            cw = (total_w - gap * (n - 1)) // n
            ph = int(cw * 1.2)
            y0 = Inches(1.85)
            for i, (ad, iller, izah, hint) in enumerate(shexsler):
                x = Inches(0.6) + i * (cw + gap)
                sekil_yeri(x, y0, cw, ph, hint)
                self.textbox(slide, ad, x, y0 + ph + Inches(0.1), cw, Inches(0.4),
                             size=14, bold=True, color=self.t["text_main"],
                             align=PP_ALIGN.CENTER)
                self.textbox(slide, iller, x, y0 + ph + Inches(0.48), cw, Inches(0.3),
                             size=11, bold=True, color=self.t["accent"],
                             align=PP_ALIGN.CENTER)
                self.textbox(slide, izah, x, y0 + ph + Inches(0.8), cw, Inches(0.7),
                             size=11, color=self.t["text_body"], align=PP_ALIGN.CENTER)
        return slide

    def saxla(self, fayl_adi="teqdimat.pptx"):
        self.prs.save(fayl_adi)
        print(f"✅ Yaradıldı: {fayl_adi} | Tema: {self.theme_name} | Slayd: {self.slide_no}")